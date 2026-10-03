"""
Copy the exercise questions out of a Tamil Nadu (Samacheer Kalvi) textbook PDF into an Excel file for review.

    python tools/extract_exercises.py BOOK.pdf --class 10 --subject social

Units come from the hidden page tag in SCERT PDFs (e.g. "01_History_Unit_1_EM.indd"), matched with the
"book" values in data/subjects.json. The kind of each question (MCQ, fill, give reasons, distinguish,
answer briefly, in detail, map ...) comes from the exercise sub-heading (I., II., III. ...).

Output: tools/review/<class>-<subject>.xlsx -> check yellow rows -> copy to tools/ -> python tools/excel_to_json.py
It only COPIES questions from the book. It never writes new ones.
"""
import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
import bookreader as br  # noqa: E402
from kinds import VALID_TYPES, classify  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]

EX_START = re.compile(r"^(?:TEXTBOOK\s+)?(?:EXERCISES?|EVALUATION)$", re.I)
EX_END = re.compile(r"^(?:REFERENCE|REFERENCES|REFERENCE BOOKS|BOOKS FOR REFERENCE|INTERNET RESOURCES?|"
                    r"WEB ?SITES?|GLOSSARY|CONCEPT MAP)\b", re.I)
ICT_LINE = re.compile(r"^(?:ICT CORNER|Steps?\b|Step\s*[–-]\s*\d|.*\bStep-\d|[•·▪●]|Website URL|URL|Through this activity|Scan the QR|"
                      r"Open the Browser|Click\b|https?://|\*Pictures)", re.I)
HEAD_RE = re.compile(r"^([IVX]{1,4})\s*[.)]?\s+(\S.*)$")
QNO_RE = re.compile(r"^(\d{1,2})\s*[.)]\s*(.*)$")
SECNUM_RE = re.compile(r"^(\d{1,2})\.\d{1,2}(?:\.\d{1,2})?\s+[A-Z]")
TAMIL_RE = re.compile(r"[\u0B80-\u0BFF]")
OPT_RE = re.compile(r"(?:(?<=\s)|^)\(?([a-d])\s?(?:\)|\.(?=\s))\s*")
ROMAN_OPT_RE = re.compile(r"(?:(?<=\s)|^)\(?(i|ii|iii|iv)\)\s*")
BLOCK_TYPES = {"match", "map"}  # whole block = one question


def split_options(text, roman=False):
    seq = ["i", "ii", "iii", "iv"] if roman else ["a", "b", "c", "d"]
    marks, k = [], 0
    for m in (ROMAN_OPT_RE if roman else OPT_RE).finditer(text):
        if m.group(1).lower() == seq[k]:
            marks.append(m)
            k += 1
            if k == 4:
                break
    if len(marks) < 2:
        return text.strip(), []
    stem = text[: marks[0].start()].strip()
    opts = [text[m.end(): (marks[i + 1].start() if i + 1 < len(marks) else len(text))].strip(" ,;")
            for i, m in enumerate(marks)]
    return stem, opts


def tidy(t):
    t = re.sub(r"\s+", " ", t).strip()
    t = re.sub(r"(?<=[.?])\s+(?:[A-Z&]{2,}\s*)+$", "", t)          # "... plains. CLIMATE AND NATURAL"
    t = re.sub(r"(fi|fl|ff) (?=[a-z])", r"\1", t)
    t = re.sub(r"(\w)- (\w)", r"\1\2", t)
    t = re.sub(r"\s+([.,?;:])", r"\1", t)
    t = re.sub(r"(_{3,})\s*(_{3,})", r"\1", t)
    return t


def norm(line):
    return re.sub(r"[\d_]+", "", line).strip()


def load_unit_map(cls, subject):
    cat = json.loads((ROOT / "data" / "subjects.json").read_text(encoding="utf-8"))["classes"][cls]
    subj = next(s for s in cat["subjects"] if s["id"] == subject)
    secs = subj.get("sections") or [{"id": "", "units": subj.get("units", [])}]
    return {str(u["book"]).lower(): (sec["id"], u["no"]) for sec in secs for u in sec["units"] if "book" in u}


def slug_key(slug, unit_map):
    s = slug.lower()
    for key in sorted(unit_map, key=len, reverse=True):
        if re.search(r"(?:^|[_\-])" + re.escape(key) + r"(?:$|[_\-])", s):
            return key
    return None


def running_titles(doc):
    c = Counter()
    for p in doc:
        for l in set(norm(t) for t, y in br.page_rows(p) if y < 0.1 or y > 0.88):
            c[l] += 1
    return {l for l, n in c.items()
            if n >= 4 and len(l) > 3 and not OPT_RE.search(l)
            and not (QNO_RE.match(l) or HEAD_RE.match(l) or EX_START.match(l) or EX_END.match(l))}


def extract(pdf, unit_map):
    doc = br.open_book(pdf)
    noise = running_titles(doc)
    out, key = [], None
    st = dict(in_ex=False, skip=False, head_i=-1, head="", qtype=None, unsure=False, shared=[], q=None, qno=0,
              collecting=False)

    def flush():
        if st["q"] and st["q"]["text"].strip():
            out.append(st["q"])
        st["q"] = None

    def new_q(text, page):
        flush()
        st["q"] = {"key": key, "type": st["qtype"], "text": text, "page": page,
                   "heading": tidy(st["head"]), "unsure": st["unsure"], "shared": st["shared"]}

    for pno in range(doc.page_count):
        page = doc[pno]
        slug = br.page_slug(page)
        lines = br.page_lines(page)
        if st["in_ex"] or any(EX_START.match(l) for l in lines):
            lines = br.page_lines(page, blanks=True)
        new_key = slug_key(slug, unit_map) if slug else None
        if not slug:
            for l in lines:
                m = SECNUM_RE.match(l)
                if m and m.group(1).lower() in unit_map:
                    new_key = m.group(1)
                    break
        if slug and not new_key and key:
            flush()
            key, st["in_ex"] = None, False
        if new_key and new_key != key:
            flush()
            key, st["in_ex"] = new_key, False
        if not key:
            continue

        for line in lines:
            line = re.sub(r"^(?:_{3,}\s*){3,}", "", line).strip()   # box borders drawn as rows of lines
            if not line.replace("_", "").strip():
                continue
            if norm(line) in noise or re.fullmatch(r"\d{1,3}", line) or TAMIL_RE.search(line):
                continue
            if EX_START.match(line):
                flush()
                st.update(in_ex=True, skip=False, head_i=-1, qtype=None, collecting=False)
                continue
            if not st["in_ex"]:
                continue
            if ICT_LINE.match(line):  # ICT corner box text: ignore the line, keep going
                if st["qtype"] in BLOCK_TYPES:
                    flush()
                continue
            if EX_END.match(line) and ":" not in line and len(line) < 40:
                flush()
                st["in_ex"] = False
                continue

            hm = HEAD_RE.match(line)
            nxt = ROMAN[st["head_i"] + 1] if st["head_i"] + 1 < len(ROMAN) else None
            if hm and hm.group(1) == nxt:
                flush()
                st.update(head_i=st["head_i"] + 1, head=hm.group(2).lstrip("_ "), skip=False, qno=0, collecting=True)
                continue
            if st["skip"] or st["head_i"] < 0:
                continue

            if st["collecting"]:
                qm = QNO_RE.match(line)
                if not qm or int(qm.group(1)) != 1:
                    st["head"] += " " + line
                    continue
                t, unsure = classify(st["head"])
                shared = []
                if t == "mcq":
                    _, shared = split_options(st["head"])
                    if len(shared) < 2:
                        _, shared = split_options(st["head"], roman=True)
                st.update(qtype=t, unsure=unsure, shared=shared, collecting=False)
                if t is None:
                    st["skip"] = True
                    continue
                if t in BLOCK_TYPES:
                    new_q(line, pno + 1)
                    continue

            if st["qtype"] in BLOCK_TYPES:
                if st["q"]:
                    last = st["q"]["text"].rstrip().split("\n")[-1]
                    # a map list ends when a non-numbered line follows a finished item ("…industries.")
                    if st["qtype"] == "map" and not QNO_RE.match(line) and last.rstrip().endswith("."):
                        flush()
                        continue
                    st["q"]["text"] += "\n" + line
                continue

            qm = QNO_RE.match(line)
            if qm and int(qm.group(1)) == st["qno"] + 1:
                st["qno"] += 1
                new_q(qm.group(2), pno + 1)
            elif st["q"]:
                if st["qtype"] == "mcq" and len(split_options(st["q"]["text"])[1]) == 4 \
                        and not OPT_RE.match(line):
                    continue
                # a finished sentence followed by a lowercase line = text from a side box, not a wrap
                if st["qtype"] not in ("mcq", "match") and re.search(r"[.?]\s*$", st["q"]["text"]) and line[:1].islower():
                    continue
                st["q"]["text"] += " " + line
    flush()
    return out, doc


def build_rows(items, cls, subject, unit_map):
    rows = []
    for it in items:
        sec, unit = unit_map[it["key"]]
        t, raw = it["type"], it["text"]
        if t in BLOCK_TYPES:
            text = "\n".join(tidy(l.replace("______", " ")) for l in raw.split("\n") if l.replace("_", "").strip())
        else:
            text = tidy(raw if t in ("fill", "mcq") else re.sub(r"\s*_{3,}\s*", " ", raw))
        if t in ("short", "long", "reason", "distinguish", "oneword"):
            # "…a mass leader. Nationalism: Previous' to see" -> drop the unfinished tail after a complete sentence
            m = re.match(r"^(.+?[.?])\s+(.*)$", text)
            if m and re.search(r"https?://|\bStep|\w+:\s|’ to ", m.group(2)) and not re.search(r"[.?]$", m.group(2).split("http")[0].strip()):
                text = m.group(1)
        opts, why = [], []
        shared = False
        if t == "mcq":
            text, opts = split_options(text)
            text = tidy(text)
            opts = [re.sub(r"[\s_]+$", "", re.split(r"[\x80-\x9f\uf000-\uf8ff]|Points to Remember", o)[0]) for o in opts]
            if not text:
                text = re.sub(r"^[\W_]+", "", it["heading"]).strip() or "Choose the correct answer"
            if not opts and it["shared"]:
                opts, shared = it["shared"], True
            if len(opts) != 4:
                why.append(f"MCQ with {len(opts)} options")
            elif not shared and max(len(o) for o in opts) > 90:
                why.append("option text too long – extra text?")
        if t == "fill" and "__" not in text:
            # the blank was a drawn line at the end of the sentence: "… in the year ______."
            text = (text[:-1] if text.endswith(".") else text).rstrip() + " ______."
            why.append("blank added at end – check position")
        if it["unsure"]:
            why.append(f"heading not recognised: {it['heading'][:40]}")
        if t == "match":
            why.append("match table – check layout")
        if len(text) < 6:
            why.append("very short")
        include = "no" if (t == "mcq" and len(opts) < 2) or not text else "yes"
        if include == "no":
            why.append("EXCLUDED until fixed")
        rows.append([int(cls), subject, sec, unit, t, text, *(opts + [""] * 4)[:4], "", "", include,
                     "YES" if why else "", "; ".join(why), it["page"], it["heading"][:80]])
    return rows


COLS = ["class", "subject", "section", "unit", "type", "question", "option_a", "option_b", "option_c", "option_d",
        "answer", "image", "include", "needs_check", "why", "book_page", "book_heading"]


def write_xlsx(rows, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "review"
    ws.append(COLS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2563EB")
    yellow = PatternFill("solid", fgColor="FEF3C7")
    for r in rows:
        ws.append(r)
        if r[13] == "YES":
            for c in ws[ws.max_row]:
                c.fill = yellow
    for i, w in enumerate([6, 9, 11, 5, 11, 70, 18, 18, 18, 18, 12, 10, 8, 10, 28, 9, 30]):
        ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = w
    ws.freeze_panes = "G2"
    dv = DataValidation(type="list", formula1='"' + ",".join(VALID_TYPES) + '"')
    ws.add_data_validation(dv)
    dv.add("E2:E9000")
    inc = DataValidation(type="list", formula1='"yes,no"')
    ws.add_data_validation(inc)
    inc.add("M2:M9000")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pdf")
    ap.add_argument("--class", dest="cls", required=True)
    ap.add_argument("--subject", required=True)
    ap.add_argument("--out")
    a = ap.parse_args()

    unit_map = load_unit_map(a.cls, a.subject)
    if not unit_map:
        sys.exit(f"❌ No 'book' keys for class {a.cls} {a.subject} in data/subjects.json")
    items, _ = extract(a.pdf, unit_map)
    if not items:
        sys.exit("❌ No exercise questions found (scanned PDF, or headings differ from EXERCISE / TEXTBOOK EVALUATION).")
    rows = build_rows(items, a.cls, a.subject, unit_map)
    out = Path(a.out) if a.out else ROOT / "tools" / "review" / f"{a.cls}-{a.subject}.xlsx"
    write_xlsx(rows, out)

    print(f"✅ {len(rows)} questions → {out}")
    for (sec, u), n in sorted(Counter((r[2], r[3]) for r in rows).items()):
        types = Counter(r[4] for r in rows if (r[2], r[3]) == (sec, u))
        print(f"   {sec or 'unit'} {u}: {n:3}  " + ", ".join(f"{t}={c}" for t, c in types.most_common()))
    missing = [k for k in unit_map if k not in {it['key'] for it in items}]
    if missing:
        print("⚠️  No questions found for:", ", ".join(missing))
    print(f"⚠️  {sum(r[13] == 'YES' for r in rows)} rows marked needs_check (yellow).")


if __name__ == "__main__":
    main()
