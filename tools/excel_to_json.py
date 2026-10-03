"""
Convert question bank Excel files into the JSON the website reads.

    python tools/excel_to_json.py            # every tools/*.xlsx  (NOT tools/review/)
    python tools/excel_to_json.py a.xlsx     # specific files

Columns (header row, any order, extra columns ignored):
    class | subject | section | unit | type | question | option_a..option_d | answer | image | include | book_heading

- type: mcq, fill, match, truefalse, oneword, short, reason, distinguish, long, timeline, problem, map
- include = no  -> row skipped
- If the row has a book_heading such as "Give reasons" or "Distinguish between", the type is set to
  reason / distinguish automatically (older Excel files had these as "short").
"""
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from kinds import VALID_TYPES, classify  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
OPTION_COLS = ["option_a", "option_b", "option_c", "option_d"]


def clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def load_catalog():
    cat = json.loads((DATA / "subjects.json").read_text(encoding="utf-8"))["classes"]
    out = {}
    for cls, c in cat.items():
        for s in c["subjects"]:
            secs = s.get("sections") or [{"id": "main", "units": s.get("units", [])}]
            out[(cls, s["id"])] = {sec["id"]: {u["no"] for u in sec["units"]} for sec in secs}
    return out


def main(files):
    catalog = load_catalog()
    groups = defaultdict(list)
    errors, fixed = [], Counter()
    for f in files:
        wb = load_workbook(f, data_only=True, read_only=True)
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                continue
            cols = {clean(h).lower(): i for i, h in enumerate(header) if h}
            if "question" not in cols:
                continue
            for r, row in enumerate(rows, start=2):
                def get(name):
                    i = cols.get(name)
                    return clean(row[i]) if i is not None and i < len(row) else ""

                if not get("question"):
                    continue
                qtype = get("type").lower()
                include = get("include").lower() not in {"no", "n", "0", "false"}

                # older files: "Give reasons" / "Distinguish between" were saved as short (or mcq)
                heading_kind, _ = classify(get("book_heading")) if get("book_heading") else (None, False)
                if heading_kind in ("reason", "distinguish") and qtype in ("short", "mcq", "long", "oneword"):
                    fixed[heading_kind] += 1
                    qtype = heading_kind
                    if not include and "EXCLUDED" in get("why"):
                        include = True
                if not include:
                    continue

                where = f"{Path(f).name} / {ws.title} / row {r}"
                cls, subj = get("class"), get("subject").lower()
                sec = get("section").lower() or "main"
                sections = catalog.get((cls, subj))
                if sections is None:
                    errors.append(f"{where}: class '{cls}' + subject '{subj}' not in subjects.json"); continue
                if sec not in sections:
                    errors.append(f"{where}: section '{sec}' not found for {subj}"); continue
                try:
                    unit = int(get("unit"))
                except ValueError:
                    errors.append(f"{where}: unit must be a number"); continue
                if unit not in sections[sec]:
                    errors.append(f"{where}: unit {unit} not listed for {subj}/{sec}"); continue
                if qtype not in VALID_TYPES:
                    errors.append(f"{where}: type '{qtype}' must be one of {VALID_TYPES}"); continue
                options = [get(c) for c in OPTION_COLS if get(c)] if qtype == "mcq" else []
                if qtype == "mcq" and len(options) < 2:
                    errors.append(f"{where}: MCQ needs at least 2 options (or set include = no)"); continue
                q = {"section": sec, "unit": unit, "type": qtype, "q": get("question")}
                if options:
                    q["options"] = options
                for extra in ("answer", "image"):
                    if get(extra):
                        q[extra] = get(extra)
                groups[(cls, subj)].append(q)

    if errors:
        print("❌ Fix these rows:\n  " + "\n  ".join(errors[:200]))
        sys.exit(1)

    for kind, n in fixed.items():
        print(f"ℹ️  {n} questions set to '{kind}' from their textbook heading")
    for (cls, subj), qs in groups.items():
        n = defaultdict(int)
        for q in qs:
            n[(q["section"], q["unit"])] += 1
            q["id"] = f"{cls}-{subj}-{q['section']}-{q['unit']}-{n[(q['section'], q['unit'])]:03d}"
        out = DATA / cls / f"{subj}.json"
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(json.dumps({"questions": qs}, ensure_ascii=False, indent=1), encoding="utf-8")
        kinds = Counter(q["type"] for q in qs)
        print(f"✅ {out.relative_to(ROOT)}: {len(qs)} questions  ({', '.join(f'{k}={v}' for k, v in kinds.most_common())})")


if __name__ == "__main__":
    args = sys.argv[1:] or sorted(str(p) for p in (ROOT / "tools").glob("*.xlsx") if not p.name.startswith("~$"))
    if not args:
        sys.exit("No .xlsx files in tools/")
    main(args)
